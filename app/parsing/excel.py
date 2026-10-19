"""添付Excel（累積型の日報ブック）の読み取り。

4院とも「月ごとのシート × 日別の行」の累積型。1つのファイルに過去の日と月合計が含まれるので、
- 日別の行だけを読み、月合計・予測・目標の行は読まない（照合用の月合計は別に返す）
- 列は見出しの文字で探す（セル番号を固定しない）。必須の見出しが見つからなければ書式変更として止める
- 数式セルは保存済みの計算結果を使い、結果が無い/古い場合は検知して要確認にする
- 抽出値には シート名・セル位置・原表記・数式 を残す
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl.utils import get_column_letter

from ..normalize import NumValue, from_cell, squash
from ..rules import EXCEL_SUMMARY_LABELS
from .formula import evaluate
from .xlsx_reader import Workbook


class ExcelFormatError(Exception):
    """必須の見出しが見つからない等、書式変更の疑い。誤ったセルから取り込まないため処理を止める。"""


@dataclass
class ExcelObs:
    target_date: date
    metric: str
    dims: dict
    num: NumValue
    sheet: str
    cell: str
    formula: str | None
    notes: list[str] = field(default_factory=list)
    definition_confirmed: bool = True

    @property
    def location(self) -> str:
        return f"{self.sheet}!{self.cell}"


@dataclass
class SheetLayout:
    sheet: str
    year: int
    month: int
    header_row: int
    columns: list[dict]          # {key, metric, dims, col, label, definition_confirmed}
    closed_col: int | None
    signature: str
    missing_optional: list[str]
    hidden_cols: list[str]


@dataclass
class DayRow:
    day: date
    row: int
    closed: bool
    has_input: bool
    obs: list[ExcelObs]


@dataclass
class WorkbookResult:
    layouts: dict[tuple[int, int], SheetLayout]
    days: dict[date, DayRow]
    month_totals: dict[tuple[int, int], dict[str, float | None]]
    warnings: list[str]
    sheet_names: list[str]

    def last_input_day(self, year: int, month: int) -> date | None:
        ds = [d for d, r in self.days.items() if d.year == year and d.month == month and r.has_input]
        return max(ds) if ds else None

    @property
    def signature(self) -> str:
        sigs = sorted({l.signature for l in self.layouts.values()})
        return hashlib.sha256("|".join(sigs).encode()).hexdigest()[:16]


# ---------------------------------------------------------------------------

def _labels(ws, row: int, max_col: int, merged: list) -> dict[int, str]:
    out = {}
    for c in range(1, max_col + 1):
        v = ws.cell(row=row, column=c).value
        if v is not None and str(v).strip():
            out[c] = squash(v)
    for rng in merged:  # 結合セルは左上の値で埋める
        if rng.min_row <= row <= rng.max_row:
            top = ws.cell(row=rng.min_row, column=rng.min_col).value
            if top is not None and str(top).strip():
                for c in range(rng.min_col, rng.max_col + 1):
                    out.setdefault(c, squash(top))
    return out


def _find_header_row(ws, header_labels: list[str]) -> int | None:
    targets = {squash(x) for x in header_labels}
    for r in range(1, 16):
        v = ws.cell(row=r, column=1).value
        if v is not None and squash(v) in targets:
            return r
    return None


def _match_cols(labels_by_rel: dict[str, dict[int, str]], rows: dict[str, list[str]]) -> list[int]:
    cols = None
    for rel, subs in rows.items():
        lab = labels_by_rel.get(rel, {})
        hit = {c for c, text in lab.items() if all(squash(s) in text for s in subs)}
        cols = hit if cols is None else cols & hit
    return sorted(cols or [])


def _layout(ws, rule: dict, year: int, month: int) -> SheetLayout:
    h = _find_header_row(ws, rule["header_labels"])
    if h is None:
        raise ExcelFormatError(f"シート「{ws.title}」に見出し行（{'/'.join(rule['header_labels'])}）が見つからない")
    max_col = ws.max_column
    merged = list(ws.merged_cells.ranges)
    labels = {str(-k) if k else "0": _labels(ws, h - k, max_col, merged) for k in range(0, 4) if h - k >= 1}
    columns: list[dict] = []
    missing_required, missing_optional = [], []
    for i, cr in enumerate(rule.get("cells", [])):
        n = cr.get("occurrence", 1)
        cols = []
        for spec in [cr["rows"], *cr.get("alt_rows", [])]:  # 書式の版ごとの見出しの候補
            cols = _match_cols(labels, spec)
            if len(cols) >= n:
                break
        key = f"{cr['metric']}{json.dumps(cr.get('dims', {}), ensure_ascii=False, sort_keys=True)}"
        if len(cols) < n:
            (missing_required if cr.get("required") else missing_optional).append(
                f"{cr['metric']}({'/'.join(sum(cr['rows'].values(), []))})")
            continue
        col = cols[n - 1]
        columns.append(dict(key=key, metric=cr["metric"], dims=dict(cr.get("dims", {})), col=col,
                            label=labels.get("0", {}).get(col) or labels.get("-1", {}).get(col, ""),
                            definition_confirmed=cr.get("definition_confirmed", True)))
    if missing_required:
        raise ExcelFormatError(f"シート「{ws.title}」で必須の見出しが見つからない: {', '.join(missing_required)}")

    # カテゴリ列（上位見出し×下位見出し×計測）
    for cat_rule in rule.get("category", []):
        g_lab = labels.get(cat_rule["group_row"], {})
        c_lab = labels.get(cat_rule["cat_row"], {})
        m_lab = labels.get(cat_rule["measure_row"], {}) if cat_rule.get("measure_row") else {}
        current_group, group_start, current_cat = None, None, None
        occ: dict[str, int] = {}
        for c in range(1, max_col + 1):
            g = g_lab.get(c)
            if g and g != current_group:
                current_group, group_start, current_cat = g, c, None
                occ[g] = occ.get(g, 0) + 1
            if current_group is None:
                continue
            spec = None
            for gs in cat_rule["groups"]:
                lab = squash(gs["label"])
                ok = current_group == lab if gs.get("exact") else lab in current_group
                if ok and occ.get(current_group, 0) == gs.get("occurrence", occ.get(current_group, 0)):
                    spec = gs
                    break
            if spec is None:
                continue
            if c_lab.get(c):
                current_cat = c_lab[c]
            cat = current_cat
            if not cat or any(x in cat for x in cat_rule.get("exclude_cat", [])):
                continue
            if cat_rule.get("measure_row"):
                metric = next((mv for mk, mv in cat_rule["measures"].items() if m_lab.get(c) == squash(mk)), None)
            else:
                metric = cat_rule.get("default_metric")
            if not metric:
                continue
            dims = {"category": cat}
            if spec.get("market"):
                dims["market"] = spec["market"]
            key = f"{metric}{json.dumps(dims, ensure_ascii=False, sort_keys=True)}"
            if any(x["key"] == key for x in columns):
                continue
            columns.append(dict(key=key, metric=metric, dims=dims, col=c, label=cat, definition_confirmed=True))

    closed_col = None
    if rule.get("closed_col"):
        cc = _match_cols(labels, rule["closed_col"])
        closed_col = cc[0] if cc else None

    sig_src = sorted((x["key"], get_column_letter(x["col"])) for x in columns)
    signature = hashlib.sha256(json.dumps(sig_src, ensure_ascii=False).encode()).hexdigest()[:16]
    hidden = [k for k, d in ws.column_dimensions.items() if d.hidden]
    return SheetLayout(ws.title, year, month, h, columns, closed_col, signature, missing_optional, hidden)


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _sheet_month(name: str, pattern: str) -> tuple[int, int] | None:
    m = re.search(pattern, squash(name))
    if not m:
        return None
    return int(m.group("y")), int(m.group("m"))


def parse_workbook(data: bytes, rule: dict, months: set[tuple[int, int]] | None = None) -> WorkbookResult:
    try:
        wb_v = Workbook(data)
    except Exception as e:  # 壊れたファイル・xls 等
        raise ExcelFormatError(f"Excelを開けない: {type(e).__name__}: {e}") from e

    layouts: dict[tuple[int, int], SheetLayout] = {}
    days: dict[date, DayRow] = {}
    totals: dict[tuple[int, int], dict[str, float | None]] = {}
    warnings: list[str] = []
    skip = [squash(s).lower() for s in rule.get("skip_sheets", [])]

    for name in wb_v.sheetnames:
        if any(s in squash(name).lower() for s in skip):
            continue
        ym = _sheet_month(name, rule["sheet_name"])
        if not ym or (months is not None and ym not in months):
            continue
        ws_v = ws_f = wb_v[name]
        if ws_v.sheet_state != "visible":
            warnings.append(f"非表示シート「{name}」")
        if ym in layouts:
            warnings.append(f"同じ月のシートが複数ある: 「{layouts[ym].sheet}」「{name}」（先のシートを使用）")
            continue
        lay = _layout(ws_v, rule, *ym)
        layouts[ym] = lay
        hidden_set = set(lay.hidden_cols)

        r = lay.header_row + 1
        while r <= ws_v.max_row:
            a = ws_v.cell(row=r, column=1).value
            d = _as_date(a) or wb_v.to_date(a)
            if d is None:
                if a is not None and any(squash(x) in squash(a) for x in EXCEL_SUMMARY_LABELS):
                    totals[ym] = {}
                    # 月合計行（例: 9ruの「予定日数」行）は照合用に保持する
                    if "予定日数" in squash(a):
                        for col in lay.columns:
                            v = ws_v.cell(row=r, column=col["col"]).value
                            totals[ym][col["key"]] = float(v) if isinstance(v, (int, float)) else None
                    break
                r += 1
                if r > lay.header_row + 45:
                    break
                continue
            if (d.year, d.month) != ym:
                warnings.append(f"シート「{name}」{r}行目の日付 {d} がシートの月と異なる")
                r += 1
                continue
            closed = False
            if lay.closed_col:
                cv = ws_v.cell(row=r, column=lay.closed_col).value
                closed = cv is not None and str(cv).strip() != ""
            obs, has_input = [], False
            for col in lay.columns:
                cell_v = ws_v.cell(row=r, column=col["col"])
                cell_f = ws_f.cell(row=r, column=col["col"])
                formula = cell_f.formula
                num = from_cell(cell_v.value)
                notes = []
                if formula is None and cell_v.value is not None:
                    has_input = True
                if formula is not None:
                    if cell_v.value is None:
                        num = NumValue("error", None, "", "数式の計算結果が保存されていない")
                        notes.append("計算結果なし（要再計算）")
                    else:
                        ev = evaluate(formula, ws_v)
                        if ev is not None and num.value is not None and abs(ev - num.value) > 0.5:
                            num.suspicious = f"保存済みの計算結果 {num.value:g} と再計算 {ev:g} が一致しない（計算結果が古い可能性）"
                            notes.append("計算結果が古い可能性")
                        if num.state == "zero" and not closed:
                            notes.append("数式の結果の0")
                letter = get_column_letter(col["col"])
                if letter in hidden_set:
                    notes.append("非表示列")
                obs.append(ExcelObs(d, col["metric"], dict(col["dims"]), num, name, f"{letter}{r}", formula,
                                    notes, col["definition_confirmed"]))
            if d in days:
                warnings.append(f"日付 {d} の行が複数ある（{days[d].row}行目と{r}行目）")
            else:
                days[d] = DayRow(d, r, closed, has_input, obs)
            r += 1

    if months and not layouts:
        raise ExcelFormatError(f"対象月のシートが見つからない（対象: {sorted(months)}、シート: {wb_v.sheetnames}）")
    return WorkbookResult(layouts, days, totals, warnings, list(wb_v.sheetnames))


def recalc_with_libreoffice(data: bytes, soffice: str) -> bytes:
    """LibreOffice（headless）で再計算して保存し直す。計算結果が無いファイル向け。"""
    if not soffice or not shutil.which(soffice) and not Path(soffice).exists():
        raise RuntimeError("LibreOffice が見つからない（LIBREOFFICE_PATH）")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in.xlsx"
        src.write_bytes(data)
        out_dir = Path(td) / "out"
        out_dir.mkdir()
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out_dir), str(src)],
                       check=True, timeout=180, capture_output=True)
        return (out_dir / "in.xlsx").read_bytes()
