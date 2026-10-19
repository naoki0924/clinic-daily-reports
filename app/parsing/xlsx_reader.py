"""軽量な xlsx 読み取り（必要なシートだけを XML から直接読む）。

openpyxl でブック全体を2回（計算結果用・数式用）読み込むと、大きな日報（約9MB・25シート）で
1件あたり数秒・数GBのメモリを使うため、対象月のシートだけを読む。
提供するもの：保存済みの計算結果、数式（共有数式は展開）、結合セル、非表示の行・列、シートの表示状態。
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from xml.etree import ElementTree as ET

from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
_REF = re.compile(r"([A-Z]+)(\d+)")
_EPOCH = datetime(1899, 12, 30)


@dataclass
class Range:
    min_col: int
    min_row: int
    max_col: int
    max_row: int


@dataclass
class _Cell:
    value: object = None
    formula: str | None = None


@dataclass
class ColDim:
    hidden: bool


@dataclass
class Sheet:
    title: str
    sheet_state: str
    cells: dict = field(default_factory=dict)       # (row, col) -> _Cell
    merged: list = field(default_factory=list)
    hidden_cols: dict = field(default_factory=dict)  # letter -> ColDim
    hidden_rows: set = field(default_factory=set)
    max_row: int = 0
    max_column: int = 0

    # openpyxl 互換の最小インターフェース
    def cell(self, row: int, column: int) -> _Cell:
        return self.cells.get((row, column)) or _Cell()

    @property
    def merged_cells(self):
        class _M:
            def __init__(self, r):
                self.ranges = r
        return _M(self.merged)

    @property
    def column_dimensions(self) -> dict:
        return self.hidden_cols


class Workbook:
    def __init__(self, data: bytes):
        self.z = zipfile.ZipFile(io.BytesIO(data))
        wb = ET.fromstring(self.z.read("xl/workbook.xml"))
        rels = ET.fromstring(self.z.read("xl/_rels/workbook.xml.rels"))
        rmap = {r.get("Id"): r.get("Target") for r in rels.findall("pr:Relationship", NS)}
        self._sheets: dict[str, tuple[str, str]] = {}
        self.sheetnames: list[str] = []
        for s in wb.find("m:sheets", NS).findall("m:sheet", NS):
            target = rmap[s.get(f"{{{NS['r']}}}id")]
            path = target.lstrip("/") if target.startswith("/") else "xl/" + target
            self._sheets[s.get("name")] = (path, s.get("state") or "visible")
            self.sheetnames.append(s.get("name"))
        self.date1904 = (wb.find("m:workbookPr", NS) is not None and
                         wb.find("m:workbookPr", NS).get("date1904") in ("1", "true"))
        self._ss: list[str] | None = None
        self._cache: dict[str, Sheet] = {}

    @property
    def shared_strings(self) -> list[str]:
        if self._ss is None:
            self._ss = []
            if "xl/sharedStrings.xml" in self.z.namelist():
                for _, el in ET.iterparse(io.BytesIO(self.z.read("xl/sharedStrings.xml"))):
                    if el.tag == f"{{{NS['m']}}}si":
                        # ふりがな（rPh）は除く
                        parts = []
                        for ch in el:
                            if ch.tag == f"{{{NS['m']}}}t":
                                parts.append(ch.text or "")
                            elif ch.tag == f"{{{NS['m']}}}r":  # リッチテキスト。rPh（ふりがな）は読まない
                                parts.extend(x.text or "" for x in ch.iter(f"{{{NS['m']}}}t"))
                        self._ss.append("".join(parts))
                        el.clear()
        return self._ss

    def __getitem__(self, name: str) -> Sheet:
        if name not in self._cache:
            self._cache[name] = self._load(name)
        return self._cache[name]

    def _load(self, name: str) -> Sheet:
        path, state = self._sheets[name]
        sh = Sheet(name, state)
        root = ET.fromstring(self.z.read(path))
        m = NS["m"]
        for col in root.iter(f"{{{m}}}col"):
            if col.get("hidden") in ("1", "true"):
                for i in range(int(col.get("min")), int(col.get("max")) + 1):
                    sh.hidden_cols[get_column_letter(i)] = ColDim(True)
        shared: dict[str, tuple[str, str]] = {}
        sd = root.find(f"{{{m}}}sheetData")
        for row in sd.findall(f"{{{m}}}row"):
            r = int(row.get("r"))
            if row.get("hidden") in ("1", "true"):
                sh.hidden_rows.add(r)
            for c in row.findall(f"{{{m}}}c"):
                ref = c.get("r")
                mm = _REF.match(ref)
                col = column_index_from_string(mm.group(1))
                t = c.get("t")
                v = c.find(f"{{{m}}}v")
                f = c.find(f"{{{m}}}f")
                value = None
                if v is not None and v.text is not None:
                    if t == "s":
                        value = self.shared_strings[int(v.text)]
                    elif t in ("str", "e"):
                        value = v.text if t == "str" else v.text  # エラー値は '#REF!' 等の文字列
                    elif t == "b":
                        value = v.text == "1"
                    else:
                        value = float(v.text)
                        if value.is_integer():
                            value = int(value)
                elif t == "inlineStr":
                    value = "".join(x.text or "" for x in c.iter(f"{{{m}}}t"))
                formula = None
                if f is not None:
                    if f.get("t") == "shared":
                        si = f.get("si")
                        if f.text:
                            shared[si] = (f.text, ref)
                            formula = "=" + f.text
                        elif si in shared:
                            src, src_ref = shared[si]
                            try:
                                formula = Translator("=" + src, origin=src_ref).translate_formula(ref)
                            except Exception:  # noqa: BLE001
                                formula = "=(shared)"
                    else:
                        formula = "=" + (f.text or "")
                if value is None and formula is None:
                    continue
                sh.cells[(r, col)] = _Cell(value, formula)
                sh.max_row = max(sh.max_row, r)
                sh.max_column = max(sh.max_column, col)
        mc = root.find(f"{{{m}}}mergeCells")
        if mc is not None:
            for x in mc.findall(f"{{{m}}}mergeCell"):
                c1, r1, c2, r2 = range_boundaries(x.get("ref"))
                sh.merged.append(Range(c1, r1, c2, r2))
        return sh

    def to_date(self, v):
        if isinstance(v, (int, float)) and 20000 < v < 80000:
            base = datetime(1904, 1, 1) if self.date1904 else _EPOCH
            return (base + timedelta(days=float(v))).date()
        return None
